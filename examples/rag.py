# /// script
# requires-python = ">=3.10"
# dependencies = ["chromadb", "pypdf", "python-docx"]
# ///
"""RAG over your own documents: ChromaDB finds the relevant passages, the local
claude-gateway answers using them.

    uv run examples/rag.py ask "Where do I work?"           # answers from examples/docs/
    uv run examples/rag.py ingest ~/notes/ other.md         # also index files outside docs/
    uv run examples/rag.py chat                            # interactive, retrieves per question
    uv run examples/rag.py list                            # what's indexed
    uv run examples/rag.py reset                           # delete the index

Files in examples/docs/ (start with docs/README.md, which is about you) are synced
automatically before every question: new or edited files are re-indexed and deleted
files are dropped, so just edit and ask. Text, Markdown, PDF and Word (.docx) files are
supported. Only passages close enough to the question are sent; if nothing matches, the
question goes to Claude on its own.

The index is stored in examples/.chroma. Embeddings come from Chroma's default local
model (all-MiniLM-L6-v2, downloaded once on first use), so documents never leave this PC
except the passages sent to the gateway.
"""

import argparse
import os
import sys
import urllib.error
from pathlib import Path

import chromadb

from ask import DIM, GATEWAY_URL, ask, claude_prompt, key_from_dotenv, style, write, you_prompt

DB_DIR = Path(__file__).resolve().parent / ".chroma"
DOCS_DIR = Path(__file__).resolve().parent / "docs"
COLLECTION = "docs"
CHUNK_CHARS = 1000
# Passages farther than this (squared L2 on normalized embeddings, 0 = identical, 4 = opposite)
# are treated as unrelated to the question and not sent.
MAX_DISTANCE = 1.5
TEXT_EXTENSIONS = {".txt", ".md", ".rst", ".csv", ".json", ".py", ".rs", ".html"}
EXTENSIONS = TEXT_EXTENSIONS | {".pdf", ".docx"}

ASSISTANT_SYSTEM = (
    "You are my personal AI assistant. Help me with whatever I ask: questions, writing, code, "
    "planning and research. When a message includes passages from my personal knowledge base, "
    "they describe me and my work, so use them to personalize your answers. Be direct and practical."
)


def collection() -> chromadb.Collection:
    client = chromadb.PersistentClient(path=str(DB_DIR))
    return client.get_or_create_collection(COLLECTION)


def split_long(para: str) -> list[str]:
    """Break a paragraph longer than CHUNK_CHARS at whitespace (PDF text often has no blank lines)."""
    pieces = []
    while len(para) > CHUNK_CHARS:
        cut = para.rfind(" ", 0, CHUNK_CHARS)
        cut = cut if cut > CHUNK_CHARS // 2 else CHUNK_CHARS
        pieces.append(para[:cut].strip())
        para = para[cut:].strip()
    return pieces + [para] if para else pieces


def chunk(text: str) -> list[str]:
    """Split on blank lines and pack paragraphs into ~CHUNK_CHARS chunks,
    repeating the last paragraph of each chunk at the start of the next for context."""
    paras = [piece for p in text.split("\n\n") if p.strip() for piece in split_long(p.strip())]
    chunks: list[str] = []
    cur: list[str] = []
    for p in paras:
        if cur and sum(len(x) for x in cur) + len(p) > CHUNK_CHARS:
            chunks.append("\n\n".join(cur))
            cur = cur[-1:] if len(cur[-1]) < CHUNK_CHARS // 2 else []
        cur.append(p)
    if cur:
        chunks.append("\n\n".join(cur))
    return chunks


def files_in(paths: list[str]) -> list[Path]:
    out = []
    for p in map(Path, paths):
        if p.is_dir():
            out += sorted(f for f in p.rglob("*") if f.is_file() and f.suffix.lower() in EXTENSIONS)
        elif p.is_file():
            out.append(p)
        else:
            print(f"skip: {p} not found", file=sys.stderr)
    return out


def extract_text(f: Path) -> str:
    """Plain text of a file: PDF pages and Word paragraphs/tables become blank-line separated text."""
    suffix = f.suffix.lower()
    if suffix == ".pdf":
        from pypdf import PdfReader

        return "\n\n".join((page.extract_text() or "") for page in PdfReader(f).pages)
    if suffix == ".docx":
        import docx

        d = docx.Document(str(f))
        parts = [p.text for p in d.paragraphs if p.text.strip()]
        for table in d.tables:
            parts += [" | ".join(c.text.strip() for c in row.cells) for row in table.rows]
        return "\n\n".join(parts)
    return f.read_text(encoding="utf-8")


def index_file(col: chromadb.Collection, f: Path) -> int:
    """(Re)index one file, replacing its old chunks. Returns the number of chunks."""
    source = str(f.resolve())
    col.delete(where={"source": source})
    try:
        text = extract_text(f)
    except Exception as e:  # unreadable PDF, not UTF-8, ...
        print(f"skip: {f}: {e}", file=sys.stderr)
        return 0
    pieces = chunk(text)
    if pieces:
        mtime = f.stat().st_mtime
        col.add(
            ids=[f"{source}#{i}" for i in range(len(pieces))],
            documents=pieces,
            metadatas=[{"source": source, "chunk": i, "mtime": mtime} for i in range(len(pieces))],
        )
    return len(pieces)


def sync_docs(col: chromadb.Collection) -> None:
    """Bring examples/docs/ and the index in line: index new or edited files, drop deleted ones."""
    if not DOCS_DIR.is_dir():
        return
    indexed: dict[str, float] = {}
    for m in col.get(include=["metadatas"])["metadatas"]:
        indexed[m["source"]] = m.get("mtime", 0)
    on_disk = {str(f.resolve()): f for f in files_in([str(DOCS_DIR)])}
    for source, f in on_disk.items():
        if indexed.get(source) != f.stat().st_mtime:
            n = index_file(col, f)
            print(f"(indexed {f.name}: {n} chunks)", file=sys.stderr)
    docs_prefix = str(DOCS_DIR.resolve()) + os.sep
    for source in indexed:
        if source.startswith(docs_prefix) and source not in on_disk:
            col.delete(where={"source": source})
            print(f"(removed {Path(source).name} from the index)", file=sys.stderr)


def cmd_ingest(paths: list[str]) -> None:
    col = collection()
    total = 0
    for f in files_in(paths or [str(DOCS_DIR)]):
        n = index_file(col, f)
        total += n
        print(f"indexed {f} ({n} chunks)")
    print(f"done: {total} chunks added, {col.count()} in the index")


def retrieve(question: str, k: int) -> list[tuple[str, str, float]]:
    """Up to k passages related to the question; empty if there are no docs or none are close."""
    col = collection()
    sync_docs(col)
    if col.count() == 0:
        return []
    res = col.query(query_texts=[question], n_results=min(k, col.count()))
    return [
        (doc, meta["source"], dist)
        for doc, meta, dist in zip(res["documents"][0], res["metadatas"][0], res["distances"][0])
        if dist <= MAX_DISTANCE
    ]


def build_question(question: str, hits: list[tuple[str, str, float]]) -> str:
    """The question plus related passages; just the question when there are none."""
    if not hits:
        return question
    docs = "\n".join(
        f'<document index="{i}" source="{Path(src).name}">\n{doc}\n</document>'
        for i, (doc, src, _) in enumerate(hits, 1)
    )
    return (
        "These passages were retrieved from my personal knowledge base because they might relate "
        "to my message. Passages about \"me\" describe me, the person writing. Use them where "
        "they help and mention the source file when you rely on one. If they aren't relevant, "
        "ignore them completely and don't mention them or the knowledge base.\n\n"
        f"<documents>\n{docs}\n</documents>\n\nMy message: {question}"
    )


def show_sources(hits: list[tuple[str, str, float]]) -> None:
    for i, (_, src, dist) in enumerate(hits, 1):
        print(style(f"  [{i}] {Path(src).name}  (distance {dist:.3f})", DIM), file=sys.stderr)


def cmd_ask(question: str, k: int, model: str, api_key: str) -> None:
    hits = retrieve(question, k)
    if hits:
        print(style("sources:", DIM), file=sys.stderr)
        show_sources(hits)
        print(file=sys.stderr)
    ask(api_key, model, [{"role": "user", "content": build_question(question, hits)}], system=ASSISTANT_SYSTEM)


def cmd_chat(k: int, model: str, api_key: str) -> None:
    print(style(f"RAG chat with {model} via {GATEWAY_URL}, {collection().count()} chunks indexed (Ctrl-D to quit)", DIM))
    # History keeps the plain questions and answers; only the current turn carries
    # retrieved documents, so the conversation doesn't grow with old context.
    history: list[dict] = []
    while True:
        try:
            question = input("\n" + you_prompt()).strip()
        except (EOFError, KeyboardInterrupt):
            print()
            return
        if not question:
            continue
        hits = retrieve(question, k)
        show_sources(hits)
        write(claude_prompt())
        try:
            answer = ask(
                api_key, model, history + [{"role": "user", "content": build_question(question, hits)}],
                system=ASSISTANT_SYSTEM,
            )
        except KeyboardInterrupt:
            write(style("\n(cancelled)\n", DIM))
            continue
        history += [{"role": "user", "content": question}, {"role": "assistant", "content": answer}]


def cmd_list() -> None:
    col = collection()
    sync_docs(col)
    metas = col.get(include=["metadatas"])["metadatas"]
    counts: dict[str, int] = {}
    for m in metas:
        counts[m["source"]] = counts.get(m["source"], 0) + 1
    for src, n in sorted(counts.items()):
        print(f"{n:4d} chunks  {src}")
    print(f"{col.count()} chunks total in {DB_DIR}")


def main() -> None:
    parser = argparse.ArgumentParser(description="RAG with ChromaDB + claude-gateway")
    parser.add_argument("--model", default="claude-opus-5-5")
    parser.add_argument("-k", type=int, default=4, help="passages to retrieve per question")
    sub = parser.add_subparsers(dest="cmd", required=True)
    sub.add_parser("ingest").add_argument("paths", nargs="*", help="default: examples/docs/")
    sub.add_parser("ask").add_argument("question", nargs="+")
    sub.add_parser("chat")
    sub.add_parser("list")
    sub.add_parser("reset")
    args = parser.parse_args()

    if args.cmd == "ingest":
        return cmd_ingest(args.paths)
    if args.cmd == "list":
        return cmd_list()
    if args.cmd == "reset":
        chromadb.PersistentClient(path=str(DB_DIR)).delete_collection(COLLECTION)
        return print("index deleted")

    api_key = os.environ.get("GATEWAY_API_KEY") or key_from_dotenv()
    if not api_key:
        sys.exit("No API key: set GATEWAY_API_KEY or GATEWAY_API_KEYS in .env")
    try:
        if args.cmd == "ask":
            cmd_ask(" ".join(args.question), args.k, args.model, api_key)
        else:
            cmd_chat(args.k, args.model, api_key)
    except urllib.error.HTTPError as e:
        sys.exit(f"Gateway error {e.code}: {e.reason}")
    except urllib.error.URLError as e:
        sys.exit(f"Cannot reach the gateway at {GATEWAY_URL} ({e.reason}). Is it running? (make status)")


if __name__ == "__main__":
    main()
